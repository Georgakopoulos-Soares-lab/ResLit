#!/usr/bin/env python3
"""
ResLit manual audit: random, tier-stratified selection of papers for two annotators
(Reviewer 1 comment 1; Reviewer 2 comment 1).

Input: the two ResLit database exports downloaded from the site. Each row is one
gene / mutation x source database x paper, with 'Validation Tier',
'Source Database' and a single 'PMID'.

Design
------
* Only ResLit-recovered entries can be audited against a source paper:
  Confirmed (literature-recovered), Supported, Candidate. Established entries
  (one reference DB, not recovered by ResLit) have no ResLit source paper and are
  excluded from the paper audit; they are covered by the database-based analyses.
* A record is one gene (or mutation) x paper, built from the rows with
  Source Database == 'Reslit', i.e. what ResLit extracted from that paper. Rows
  from CARD / ResFinder / NCBI are used only to set the entity's tier.
* Each paper gets ONE stratum: the least-supported tier among its records
  (Candidate < Supported < Confirmed) and the entity type (gene / mutation) of
  those records. Papers are drawn uniformly at random within each stratum (fixed
  seed), plus a simple random draw from the remaining papers. Nobody chooses
  individual papers. Each paper's inclusion probability (stratum draw + random
  draw) and weight are saved.
* Papers are split between annotators A and B (default 30 each, 10 shared for
  Cohen's kappa). Each annotator gets ONE file, *_step1_blind.xlsx: the PMIDs, to
  be filled with every resistance gene / mutation the paper's own results report
  (type, gene, nucleotide change, protein change, organism, resistance), WITHOUT
  seeing ResLit's output.
* All ResLit records of the selected papers are written with the same columns
  (*_reslit_records.tsv). reslit_audit_compare.py then compares them with the
  annotators' lists computationally (TP / FP / FN, organism and resistance agreement,
  Cohen's kappa).

Usage
-----
  python reslit_audit_select.py \
      --genes genes_export.csv --mutations mutations_export.csv \
      --n-papers 50 --overlap 10 --seed 2026 --out audit

  Run ONCE with the seed you will report. Do not rerun with other seeds.
"""
import argparse
import re
import sys

import numpy as np
import pandas as pd

AUDITABLE = ["Candidate", "Supported", "Confirmed"]          # rarest first
RANK = {t: i for i, t in enumerate(AUDITABLE)}

# Allocation of papers to strata for --n-papers 50 (scaled for other totals).
ALLOC = {
    ("Candidate", "gene"): 13,
    ("Candidate", "mutation"): 13,
    ("Supported", "gene"): 5,
    ("Supported", "mutation"): 5,
    ("Confirmed", "gene"): 4,
    ("Confirmed", "mutation"): 4,
    ("Random", "any"): 6,
}


# ----------------------------------------------------------------------------- io
PRIME_RE = re.compile(r'(\(\d+)"{1,2}(?=[)\-A-Za-z])')


def read_table(path):
    """Read xlsx/csv/tsv. Delimiter is detected from the header line. Rows with too many
    fields (an unquoted separator inside a text field) are repaired by joining the extra
    fields into the free-text column where they most likely belong, and reported."""
    import csv
    if path.lower().endswith((".xlsx", ".xls")):
        return pd.read_excel(path, dtype=str).fillna("")
    import io
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as fh:
        text = fh.read()
    # A double quote used as a double prime in gene names, e.g. aph(3")-Ib or aph(3"-Ib,
    # breaks CSV quoting. Replace it with two apostrophes: aph(3'')-Ib. Only quotes that
    # directly follow '(' + digits and precede ')', '-' or a letter are touched.
    text, n_prime = PRIME_RE.subn(lambda m: m.group(1) + "''", text)
    if n_prime:
        print(f"NOTE: {path}: fixed {n_prime} double-quote prime mark(s) in gene names "
              "while reading (input file unchanged)", file=sys.stderr)
    buf = io.StringIO(text, newline="")
    del text
    with buf as fh:
        head = fh.readline()
        sep = max(["\t", ",", ";"], key=head.count)
        fh.seek(0)
        csv.field_size_limit(10**9)
        reader = csv.reader(fh, delimiter=sep)
        header = next(reader)
        n = len(header)
        lower = [h.strip().lower() for h in header]
        # free-text columns, in order of preference, that may hold a stray separator
        text_cols = [i for name in ("key findings", "comments", "notes", "paper title",
                                    "effect", "confirmation reason")
                     for i, h in enumerate(lower) if h == name]
        rows, fixed, suspect = [], [], []
        for line_no, r in enumerate(reader, start=2):
            if len(r) == n:
                rows.append(r)
                continue
            if len(r) < n:
                rows.append(r + [""] * (n - len(r)))
                fixed.append(line_no)
                continue
            extra = len(r) - n
            k = text_cols[0] if text_cols else n - 1
            repaired = r[:k] + [sep.join(r[k:k + extra + 1])] + r[k + extra + 1:]
            rows.append(repaired)
            fixed.append(line_no)
            tier_i = lower.index("validation tier") if "validation tier" in lower else None
            pmid_i = lower.index("pmid") if "pmid" in lower else None
            ok_tier = tier_i is None or repaired[tier_i].strip().capitalize() in (
                "Confirmed", "Established", "Supported", "Candidate", "")
            ok_pmid = pmid_i is None or repaired[pmid_i].strip() == "" or \
                re.search(r"\d{4,9}", repaired[pmid_i]) is not None
            if not (ok_tier and ok_pmid):
                suspect.append(line_no)
    if fixed:
        print(f"NOTE: {path}: repaired {len(fixed)} malformed line(s) "
              f"(first: {fixed[:10]})", file=sys.stderr)
    if suspect:
        print(f"WARNING: {path}: {len(suspect)} repaired line(s) still look misaligned "
              f"(tier/PMID column): {suspect[:20]}. Check these lines in the file.",
              file=sys.stderr)
    df = pd.DataFrame(rows, columns=header, dtype=str)
    del rows
    return df


def col(df, *names):
    """First existing column among names (case/space/underscore-insensitive), else ''."""
    norm = {re.sub(r"[\s_]+", "", c).lower(): c for c in df.columns}
    for n in names:
        c = norm.get(re.sub(r"[\s_]+", "", n).lower())
        if c is not None:
            return df[c].astype(str).str.strip()
    return pd.Series([""] * len(df), index=df.index)


def split_pmids(s):
    """'123; 456 | 789' -> ['123','456','789']; drops 'Unpublished', ISBNs, etc."""
    return [t for t in re.split(r"[;,|\s]+", str(s)) if re.fullmatch(r"\d{4,9}", t)]


# ------------------------------------------------------------------- entity table
def load_export(path, kind):
    import gc
    df = read_table(path)
    gc.collect()
    tier = col(df, "Validation Tier", "Tier").str.strip().str.capitalize()
    gene = col(df, "Gene Name", "Gene")
    empty = pd.Series([""] * len(df), index=df.index)
    allele = col(df, "Allele") if kind == "gene" else empty
    prot = col(df, "Protein Change") if kind == "mutation" else empty
    nuc = col(df, "Nucleotide Change") if kind == "mutation" else empty
    if kind == "gene":
        # the site lower-cases some gene names; show the allele's spelling when only case differs
        same = allele.str.lower() == gene.str.lower()
        label = pd.Series(np.where(allele == "", gene,
                                   np.where(same, allele, gene + " [" + allele + "]")),
                          index=df.index)
        org = col(df, "Organisms Tested In", "Organism")
        key = gene.str.lower().str.replace(r"\s+", "", regex=True)
    else:
        change = pd.Series(np.where(prot != "", "p." + prot,
                                    np.where(nuc != "", "c." + nuc,
                                             col(df, "Mutation Name"))), index=df.index)
        label = gene + " " + change
        org = col(df, "Organisms Observed In", "Organism")
        key = (gene.str.lower().str.replace(r"\s+", "", regex=True) + "|" +
               change.str.lower().str.replace(r"\s+", "", regex=True))
    return pd.DataFrame({
        "entity_type": kind,
        "key": key,
        "entity": label.str.strip(),
        "Tier": tier,
        "gene": gene,
        "allele": allele,
        "nucleotide_change": nuc,
        "protein_change": prot,
        "organism": org,
        "resistance": col(df, "Confers Resistance To", "Resistance"),
        "Source": col(df, "Source Database"),
        "PMID_raw": col(df, "PMID"),
    })


def join_unique(s):
    """Distinct non-empty values, in order, joined with ' | '."""
    return " | ".join(dict.fromkeys(v for v in s if v))


def scale_alloc(n):
    tot = sum(ALLOC.values())
    raw = {k: v * n / tot for k, v in ALLOC.items()}
    alloc = {k: int(np.floor(v)) for k, v in raw.items()}
    for k in sorted(raw, key=lambda k: raw[k] - alloc[k], reverse=True)[: n - sum(alloc.values())]:
        alloc[k] += 1
    return alloc


def pubmed(p):
    return f"https://pubmed.ncbi.nlm.nih.gov/{p}/"



# ------------------------------------------------------------- step-1 workbook
STEP1_COLS = [   # (column, width, drop-down values)
    ("PMID", 11, None),
    ("PubMed", 40, None),
    ("type", 11, ["gene", "mutation"]),
    ("gene", 20, None),
    ("nucleotide_change", 18, None),
    ("protein_change", 16, None),
    ("organism", 30, None),
    ("resistance", 30, None),
]


def write_step1(path, pmids, ann):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "determinants"
    ws.append([c[0] for c in STEP1_COLS])
    for j, (name, width, _) in enumerate(STEP1_COLS, 1):
        cell = ws.cell(row=1, column=j)
        cell.fill, cell.font = PatternFill("solid", fgColor="1F4E78"), Font(bold=True, color="FFFFFF")
        ws.column_dimensions[get_column_letter(j)].width = width
    grey, band = PatternFill("solid", fgColor="E7E6E6"), PatternFill("solid", fgColor="F3F7FB")
    for i, p in enumerate(pmids):
        ws.append([p, pubmed(p)] + [""] * (len(STEP1_COLS) - 2))
        for j in range(1, len(STEP1_COLS) + 1):
            if j <= 2 or i % 2:
                ws.cell(row=ws.max_row, column=j).fill = grey if j <= 2 else band
    n = ws.max_row
    for j, (_, _, allowed) in enumerate(STEP1_COLS, 1):
        if allowed:
            dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            L = get_column_letter(j)
            dv.add(f"{L}2:{L}{n + 1000}")   # +1000 so added rows validate too
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(STEP1_COLS))}{n}"

    ws0 = wb.create_sheet("instructions")
    for line in [
        f"ResLit audit - STEP 1 (blind) - annotator {ann}",
        "",
        "For each paper, list every resistance gene or mutation that the paper's OWN results "
        "(text, tables, figures, supplement) link to resistance. Skip ones only cited as background.",
        "One row per gene/mutation x organism. Each paper has one row; for more, insert rows "
        "below it and copy the PMID. Add as many as the paper needs.",
        "type: gene (acquired resistance gene) or mutation. For a gene, write the allele if given "
        "(e.g. blaCTX-M-15). For a mutation, write the gene (e.g. gyrA) and the change.",
        "resistance: antibiotic(s) or class; separate several with ';'.",
        "If a paper reports none, leave its row empty.",
        "Do not look up these papers in ResLit. Do not discuss the shared papers with the "
        "other annotator until both of you have finished.",
    ]:
        ws0.append([line])
    ws0.column_dimensions["A"].width = 120
    ws0["A1"].font = Font(bold=True, size=13)
    wb.save(path)


# --------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--genes", required=True, help="gene export (tsv/csv/xlsx)")
    ap.add_argument("--mutations", required=True, help="mutation export (tsv/csv/xlsx)")
    ap.add_argument("--n-papers", type=int, default=50)
    ap.add_argument("--overlap", type=int, default=10, help="papers seen by both annotators")
    ap.add_argument("--seed", type=int, default=2026)
    ap.add_argument("--out", default="audit")
    a = ap.parse_args()
    rng = np.random.default_rng(a.seed)

    print("Reading input files ...", file=sys.stderr, flush=True)
    ent = pd.concat([load_export(a.genes, "gene"), load_export(a.mutations, "mutation")],
                    ignore_index=True)

    ent["row"] = np.arange(len(ent))

    # Each row keeps the tier the site computed for it. The key below is case-insensitive,
    # so it can merge site entities with different tiers (e.g. CARD 'AcrE' and ResLit 'acrE');
    # for this summary table only, such a merged entity is shown with its strongest tier.
    prio = {"Confirmed": 0, "Established": 1, "Supported": 2, "Candidate": 3}
    ent["_p"] = ent.Tier.map(prio).fillna(9)
    uk = ent.sort_values("_p", kind="stable").drop_duplicates(["entity_type", "key"])
    n_conf = int((ent.groupby(["entity_type", "key"]).Tier.nunique() > 1).sum())
    print("== Rows in the export files, by tier ==")
    print(pd.crosstab(ent.Tier, ent.entity_type, margins=True), "\n")
    print("== UNIQUE entities per tier (case-insensitive names; approximate) ==")
    print(pd.crosstab(uk.Tier, uk.entity_type, margins=True))
    if n_conf:
        print(f"   ({n_conf} entities have rows with different tiers; the strongest is shown here, "
              "the row's own tier is used for the audit)")
    print()

    # entities whose gene name holds several genes (e.g. "aph(3'')-Ib, aph(6)-Id")
    gname = uk.entity.str.replace(r"\s*\[.*\]$", "", regex=True)
    merged = uk[((uk.entity_type == "gene") & gname.str.contains(r"[,;/]\s*\S", regex=True))
                | ((uk.entity_type == "mutation") &
                   uk.entity.str.split(" ").str[0].str.contains(r"[,;]", regex=True))]
    if len(merged):
        merged[["entity_type", "entity", "Tier"]].to_csv(
            f"{a.out}_merged_gene_names.tsv", sep="\t", index=False)
        print(f"NOTE: {len(merged)} unique entities have several genes in one name "
              f"(by tier: {merged.Tier.value_counts().to_dict()}). Kept as they are for the "
              f"audit; list written to {a.out}_merged_gene_names.tsv. '/'-names may be "
              "legitimate CARD names (e.g. AAC6_30/AAC6_Ib) - review the list.\n",
              file=sys.stderr)

    # ---- records: one per entity x paper, from the rows ResLit extracted
    detail_cols = ["entity", "gene", "allele", "nucleotide_change", "protein_change",
                   "organism", "resistance"]
    is_reslit = ent.Source.str.lower().eq("reslit")
    print(f"ResLit rows in the exports: {int(is_reslit.sum()):,} of {len(ent):,} "
          f"(tiers: {ent[is_reslit].Tier.value_counts().to_dict()})\n")
    # strongest tier first, so 'first' below gives each record its strongest row tier
    aud = ent[is_reslit & ent.Tier.isin(AUDITABLE)].sort_values("_p", kind="stable")
    aud["PMID"] = aud.PMID_raw.map(split_pmids)
    n_multi = int((aud.PMID.map(len) > 1).sum())
    if n_multi:
        print(f"WARNING: {n_multi} ResLit rows list several PMIDs; each PMID is treated as a "
              "separate record. Expected one PMID per row - check the export.", file=sys.stderr)
    aud = aud.explode("PMID").dropna(subset=["PMID"])
    aud["PMID"] = aud.PMID.astype(np.int64)
    pairs = (aud.groupby(["entity_type", "key", "PMID"], sort=False)
                .agg(Tier=("Tier", "first"), n_rows=("row", "size"),
                     **{c: (c, join_unique) for c in detail_cols})
                .reset_index())
    pairs["key_id"] = pd.factorize(pairs.entity_type + "\t" + pairs.key)[0]
    del aud

    print(f"Linking entities to papers: {len(pairs):,} entity-paper records, "
          f"{pairs.PMID.nunique():,} papers ...", file=sys.stderr, flush=True)
    pairs["tier_rank"] = pairs.Tier.map(RANK).astype(int)
    best = pairs.groupby("PMID").tier_rank.min()
    pairs = pairs.join(best.rename("paper_rank"), on="PMID")
    rare = pairs[pairs.tier_rank == pairs.paper_rank]
    papers = pd.DataFrame({
        "stratum_tier": best.map(dict(enumerate(AUDITABLE))),
        "entity_type": (rare.groupby(["PMID", "entity_type"]).size().rename("n")
                        .reset_index()
                        .sort_values(["PMID", "n", "entity_type"], ascending=[True, False, True])
                        .drop_duplicates("PMID").set_index("PMID").entity_type),
        "n_records_total": pairs.groupby("PMID").size(),
    })
    print("== Sampling frame: papers per stratum ==")
    print(papers.groupby(["stratum_tier", "entity_type"]).size().to_string(), "\n")

    alloc = scale_alloc(a.n_papers)
    strata = [k for k in alloc if k[0] != "Random"]
    frame = {k: sorted(papers[(papers.stratum_tier == k[0]) & (papers.entity_type == k[1])].index)
             for k in strata}
    carry = 0
    for k in strata:
        want = alloc[k] + carry
        alloc[k] = min(want, len(frame[k]))
        carry = want - alloc[k]
        if carry:
            print(f"NOTE: stratum {k} has only {len(frame[k])} papers; {carry} slot(s) moved on",
                  file=sys.stderr)
    alloc[("Random", "any")] += carry

    rows, chosen = [], set()
    for k in strata:
        if alloc[k]:
            for p in rng.choice(frame[k], size=alloc[k], replace=False):
                rows.append((p, f"{k[0]}-{k[1]}"))
                chosen.add(p)
    pool = sorted(p for p in papers.index if p not in chosen)
    kr = min(alloc[("Random", "any")], len(pool))
    for p in rng.choice(pool, size=kr, replace=False):
        rows.append((p, "Random"))
    sel = pd.DataFrame(rows, columns=["PMID", "sampling_stratum"])
    sel = sel.join(papers, on="PMID")
    # Inclusion probability of any paper in stratum k: drawn in its stratum (n_k / N_k), or
    # not drawn there and then drawn in the random step from the remaining M papers.
    M = len(pool)
    prob = {k: (alloc[k] / len(frame[k]) if frame[k] else 0.0) for k in strata}
    prob = {k: v + (1 - v) * (kr / M if M else 0.0) for k, v in prob.items()}
    sel["inclusion_prob"] = [prob[(t, e)] for t, e in zip(sel.stratum_tier, sel.entity_type)]
    sel["weight"] = 1 / sel.inclusion_prob
    sel["PubMed"] = sel.PMID.map(pubmed)

    # annotators: shared papers round-robin across strata; the rest alternate A/B
    order = sel.sample(frac=1, random_state=a.seed)
    order["_i"] = order.groupby("sampling_stratum").cumcount()
    order = order.sort_values(["_i", "sampling_stratum"], kind="stable")
    shared = list(order.PMID[: a.overlap])
    rest = order[~order.PMID.isin(shared)].sort_values(["sampling_stratum", "_i"], kind="stable")
    assign = {p: "BOTH" for p in shared}
    for j, p in enumerate(rest.PMID):
        assign[p] = "A" if j % 2 == 0 else "B"
    sel["assigned_to"] = sel.PMID.map(assign)

    rec = pairs[pairs.PMID.isin(sel.PMID)].copy()
    rec = rec.merge(sel[["PMID", "sampling_stratum", "assigned_to"]], on="PMID")
    rec = rec.sort_values(["PMID", "tier_rank", "entity"]).reset_index(drop=True)
    rec["record_id"] = [f"R{i:05d}" for i in range(len(rec))]

    # ---- how many ResLit entries each selected paper should give
    per = (rec.pivot_table(index="PMID", columns="Tier", values="key_id", aggfunc="size",
                           fill_value=0)
              .reindex(columns=AUDITABLE, fill_value=0))
    per["genes"] = rec[rec.entity_type == "gene"].groupby("PMID").size()
    per["mutations"] = rec[rec.entity_type == "mutation"].groupby("PMID").size()
    per = per.fillna(0).astype(int)
    per["records"] = rec.groupby("PMID").size()
    per = sel.set_index("PMID")[["sampling_stratum", "assigned_to"]].join(per)
    per = per.sort_values(["assigned_to", "sampling_stratum"])
    sel = sel.join(per[AUDITABLE + ["genes", "mutations", "records"]], on="PMID")
    pd.set_option("display.width", 200)
    print("== ResLit entries per selected paper ==")
    print(per.to_string(), "\n")
    print("== Totals per annotator (ResLit records in each annotator's papers) ==")
    for ann in ("A", "B"):
        m = per[per.assigned_to.isin([ann, "BOTH"])]
        print(f"  Annotator {ann}: {len(m)} papers, {int(m.records.sum())} ResLit records "
              f"(Confirmed {int(m.Confirmed.sum())}, Supported {int(m.Supported.sum())}, "
              f"Candidate {int(m.Candidate.sum())}; genes {int(m.genes.sum())}, "
              f"mutations {int(m.mutations.sum())})")
    sh = per[per.assigned_to == "BOTH"]
    print(f"  Shared (kappa) subset: {len(sh)} papers, {int(sh.records.sum())} ResLit records")
    print(f"  All {len(per)} papers: {int(per.records.sum())} ResLit records\n")

    print(f"== Selected {sel.PMID.nunique()} papers ==")
    print(pd.crosstab(sel.sampling_stratum, sel.assigned_to, margins=True), "\n")
    print("== ResLit records to compare (tier x type) ==")
    print(pd.crosstab(rec.Tier, rec.entity_type, margins=True), "\n")
    n_cand = int((rec.Tier == "Candidate").sum())
    if n_cand < 100:
        print(f"NOTE: {n_cand} Candidate records (target >=100). Raise the Candidate allocation "
              "in ALLOC for a narrower Candidate confidence interval.", file=sys.stderr)

    print("Writing Excel files ...", file=sys.stderr, flush=True)
    # ------------------------------------------------------------------ outputs
    # ResLit records, with the same basic columns as the annotators' step-1 sheet
    rec_out = rec[["record_id", "PMID", "assigned_to", "sampling_stratum", "Tier",
                   "entity_type", "gene", "allele", "nucleotide_change", "protein_change",
                   "organism", "resistance"]].rename(columns={"entity_type": "type"})
    rec_out.insert(2, "PubMed", rec_out.PMID.map(pubmed))

    for ann in ("A", "B"):
        mine = sel[sel.assigned_to.isin([ann, "BOTH"])].sort_values("PMID")
        write_step1(f"{a.out}_annotator_{ann}_step1_blind.xlsx", list(mine.PMID), ann)
        print(f"Wrote {a.out}_annotator_{ann}_step1_blind.xlsx ({len(mine)} papers)")
    rec_out.to_csv(f"{a.out}_reslit_records.tsv", sep="\t", index=False)
    print(f"Wrote {a.out}_reslit_records.tsv ({len(rec_out)} records; NOT for the annotators)")

    # master workbook: everything in one place (for the PI / for computing metrics)
    with pd.ExcelWriter(f"{a.out}_MASTER_{len(sel)}papers.xlsx") as xw:
        sel.to_excel(xw, sheet_name="papers", index=False)
        rec_out.to_excel(xw, sheet_name="reslit_records", index=False)
        pd.crosstab(rec.Tier, rec.entity_type, margins=True).to_excel(xw, sheet_name="summary")
    print(f"Wrote {a.out}_MASTER_{len(sel)}papers.xlsx")
    sel.to_csv(f"{a.out}_selected_papers.tsv", sep="\t", index=False)
    with open(f"{a.out}_selection_log.txt", "w") as fh:
        fh.write(f"seed={a.seed} n_papers={a.n_papers} overlap={a.overlap}\n")
        fh.write(f"genes={a.genes} mutations={a.mutations}\n")
        fh.write(f"final_allocation={alloc}\n")
    print(f"Wrote {a.out}_selected_papers.tsv and {a.out}_selection_log.txt")


if __name__ == "__main__":
    main()

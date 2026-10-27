#!/usr/bin/env python3
"""Compare an annotator's genes with ResLit's alleles (type = gene only).

Within each paper, the annotator's `gene` column is compared with ResLit's
`allele` column after normalising both names: a trailing variant number such
as "_1" is dropped (blaTEM-1B_1 -> blaTEM-1B), then lower case, and every
character that is not a letter or digit removed (spaces, (), -, ', _, ...).
ResLit alleles listed together ("blaKPC-2 | blaKPC-3") are split into
separate entries. Duplicates within a paper (same gene in several organisms)
are counted once.

Aliases: an annotator name with a synonym in brackets after a space, e.g.
"IniA (Rv0342)", matches on the full name, on "IniA" or on "Rv0342".

TP = ResLit allele found in the annotator's list (precision)
FP = ResLit allele not found
FN = annotator gene not found in ResLit (recall)

Usage:
    python3 scripts/gene_allele_compare.py --letter A
    python3 scripts/gene_allele_compare.py --letter B
"""

import argparse
import math
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA = ROOT / "data"
RESULTS = ROOT / "results"
TIERS = ["Candidate", "Supported", "Confirmed"]
FILLS = {"match": "C6EFCE", "only ResLit (FP)": "FFC7CE", "only annotator (FN)": "FFEB9C"}


def norm(name):
    # drop a trailing variant number such as "_1" (blaTEM-1B_1 -> blaTEM-1B)
    name = re.sub(r"_\d+$", "", str(name).strip())
    return re.sub(r"[^0-9a-z]", "", name.lower())


# Missing "bla" prefix on beta-lactamase names (CTX-M-15 -> blaCTX-M-15), as in the
# ResLit pipeline (05_harmonised_pipeline/reslit/genes/qwen3_to_csv.py, _BLA_FAMILIES /
# _BLA_RE), plus PSE and CARB, which the pipeline list does not have.
_BLA_FAMILIES = (
    "CTX-M", "TEM", "SHV", "OXA", "NDM", "KPC", "VIM", "IMP", "GES",
    "CMY", "DHA", "FOX", "MOX", "ACT", "MIR", "ACC", "ADC", "PDC",
    "PER", "VEB", "BEL", "GIM", "SPM", "AIM", "IMI", "CAR",
    "CFE", "FRI", "ROB", "SCO", "SME", "CEP",
    "LAP", "LEN", "OKP", "OXY", "OCH",
    "PSE", "CARB",
)
_BLA_RE = re.compile(
    r"^(?!bla)(" + "|".join(re.escape(f) for f in _BLA_FAMILIES) + r")(?=[-_\d\s]|$)",
    re.IGNORECASE,
)


def add_bla(name):
    """Prefix "bla" to a beta-lactamase name written without it (annotator side)."""
    s = str(name).strip()
    return "bla" + s if _BLA_RE.match(s) else s


def annotator_norm(name):
    """Normalised annotator gene name: missing bla prefix added, then norm()."""
    return norm(add_bla(name))


def aliases(name):
    """Full name plus, for 'X (Y)', the parts X and Y (annotator side, bla prefix added)."""
    out = {annotator_norm(name)}
    m = re.match(r"^(.*?)\s+\((.*)\)\s*$", str(name).strip())
    if m:
        out.add(annotator_norm(m.group(1)))
        parts = re.split(r"[;/,]", m.group(2))
        out.update(annotator_norm(re.sub(r"^\s*(including|incl\.?|also)\s+", "", p)) for p in parts)
    return {a for a in out if a}


def read_annotator(path):
    """The annotator's determinants sheet, plus the free-text note column (10th column) as `note`."""
    raw = pd.read_excel(path, sheet_name="determinants", dtype=str)
    raw["PMID"] = raw["PMID"].str.strip()
    raw["note"] = raw.iloc[:, 9] if raw.shape[1] > 9 else None
    no_det = set(raw.loc[raw.iloc[:, 8].fillna("").str.strip().str.upper().eq("Y"), "PMID"])
    return raw, sorted(set(raw["PMID"].dropna())), no_det


def join_set(s):
    return " ; ".join(sorted(set(s.dropna())))


def wilson(k, n, z=1.96):
    if n == 0:
        return float("nan"), float("nan")
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return c - h, c + h


def rate(metric, subset, k, n):
    lo, hi = wilson(k, n)
    return {"metric": metric, "subset": subset, "k": k, "n": n,
            "value": k / n if n else float("nan"), "ci95_low": lo, "ci95_high": hi}


def joined(s):
    return " | ".join(s.dropna())


def write_workbook(out, sheets, color_sheet="side_by_side", fills=FILLS):
    """Write the sheets with bold headers, column widths and rows of `color_sheet`
    coloured by their status (2nd column)."""
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2" if name in ("summary", "method") else "B2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for j, col in enumerate(df.columns, 1):
                w = max([len(str(col))] + [len(str(v)) for v in df[col]])
                ws.column_dimensions[get_column_letter(j)].width = min(w + 2, 70)
        for row in (xw.sheets[color_sheet].iter_rows(min_row=2) if color_sheet in xw.sheets else []):
            if row[1].value in fills:
                fill = PatternFill("solid", fgColor=fills[row[1].value])
                for c in row:
                    c.fill = fill
        if "summary" in xw.sheets:
            for row in xw.sheets["summary"].iter_rows(min_row=2, min_col=5, max_col=7):
                for c in row:
                    c.number_format = "0.000"


def annotator_genes(raw):
    """Annotator gene rows -> one entry per (paper, normalised name), with aliases."""
    ann = raw[raw["type"].str.strip().str.lower().eq("gene") & raw["gene"].notna()].copy()
    ann["norm"] = ann["gene"].map(annotator_norm)
    ann = ann[ann["norm"].ne("")]
    hum = (ann.groupby(["PMID", "norm"], as_index=False)
              .agg(annotator_gene=("gene", join_set), annotator_organism=("organism", join_set),
                   annotator_note=("note", join_set)))
    hum["aliases"] = hum["annotator_gene"].map(lambda g: set().union(*(aliases(x) for x in g.split(" ; "))))
    return hum


def reslit_genes(res):
    """ResLit gene records -> one entry per (paper, normalised allele)."""
    res = res.copy()
    res["allele_one"] = res["allele"].fillna(res["gene"]).str.split("|")
    res = res.explode("allele_one")
    res["allele_one"] = res["allele_one"].str.strip()
    res["norm"] = res["allele_one"].map(norm)
    res = res[res["norm"].ne("")]
    rl = (res.sort_values("Tier", key=lambda s: s.map(TIERS.index))
             .groupby(["PMID", "norm"], as_index=False)
             .agg(reslit_gene=("gene", lambda s: " ; ".join(sorted(set(s)))),
                  reslit_allele=("allele_one", lambda s: " ; ".join(sorted(set(s)))),
                  reslit_allele_full=("allele", lambda s: " ; ".join(sorted(set(s.dropna())))),
                  reslit_tier=("Tier", "first"),
                  reslit_organism=("organism", lambda s: " ; ".join(sorted(set(s.dropna())))),
                  record_id=("record_id", lambda s: ",".join(sorted(set(s))))))
    return rl


def gene_match(n, h):
    """ResLit allele (normalised) vs one annotator entry h."""
    return n in h.aliases


def genes_alike(h1, h2):
    """Two annotator entries name the same gene."""
    return bool(h1.aliases & h2.aliases)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--annotator", default=None, help="default: data/annotator_<letter>.xlsx")
    ap.add_argument("--letter", default="A", help="annotator letter in assigned_to (A or B)")
    ap.add_argument("--reslit", default=DATA / "reslit_records.tsv")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    args.annotator = args.annotator or DATA / f"annotator_{args.letter}.xlsx"
    out = Path(args.out or RESULTS / f"gene_results_{args.letter}.xlsx")

    # Annotator: gene rows only, one entry per (paper, normalised name)
    raw, papers, no_det = read_annotator(args.annotator)
    hum = annotator_genes(raw)

    # ResLit: gene records of this annotator's papers, one entry per allele
    res = pd.read_csv(args.reslit, sep="\t", dtype=str)
    res = res[res["type"].eq("gene") & res["assigned_to"].isin([args.letter, "BOTH"]) & res["PMID"].isin(papers)]
    rl = reslit_genes(res)

    # Matching (with aliases)
    hum_alias = hum.groupby("PMID")["aliases"].apply(lambda s: set().union(*s)).to_dict()
    res_names = rl.groupby("PMID")["norm"].apply(set).to_dict()
    rl["TP"] = [n in hum_alias.get(p, set()) for p, n in zip(rl.PMID, rl.norm)]
    hum["matched"] = [bool(a & res_names.get(p, set())) for p, a in zip(hum.PMID, hum.aliases)]

    # Side by side: each ResLit allele next to the annotator gene(s) it matches
    rows = []
    for r in rl.itertuples():
        hits = hum[(hum.PMID == r.PMID) & hum.aliases.map(lambda a: r.norm in a)]
        rows.append({"PMID": r.PMID, "status": "match" if r.TP else "only ResLit (FP)",
                     "reslit_gene": r.reslit_gene, "reslit_allele": r.reslit_allele,
                     "reslit_allele_full": r.reslit_allele_full, "reslit_tier": r.reslit_tier,
                     "annotator_gene": " ; ".join(hits.annotator_gene) or None,
                     "reslit_organism": r.reslit_organism,
                     "annotator_organism": " ; ".join(hits.annotator_organism) or None,
                     "normalised_name": r.norm,
                     "annotator_note": " ; ".join(n for n in hits.annotator_note if n) or None})
    for h in hum[~hum.matched].itertuples():
        rows.append({"PMID": h.PMID, "status": "only annotator (FN)", "annotator_gene": h.annotator_gene,
                     "annotator_organism": h.annotator_organism, "normalised_name": h.norm,
                     "annotator_note": h.annotator_note or None})
    side = pd.DataFrame(rows)
    side["_o"] = side.status.map(list(FILLS).index)
    side = side.sort_values(["PMID", "_o", "normalised_name"]).drop(columns="_o")

    # Summary
    tp, n_r = int(rl.TP.sum()), len(rl)
    tp_h, n_h = int(hum.matched.sum()), len(hum)
    prec, rec = rate("precision", "all", tp, n_r), rate("recall", "all", tp_h, n_h)
    f1 = 2 * prec["value"] * rec["value"] / (prec["value"] + rec["value"]) if tp and tp_h else 0.0
    summary = pd.DataFrame([
        prec,
        *[rate("precision", t, int(rl.loc[rl.reslit_tier == t, "TP"].sum()), int((rl.reslit_tier == t).sum()))
          for t in TIERS if (rl.reslit_tier == t).any()],
        rec,
        {"metric": "F1", "subset": "all", "value": f1},
        {"metric": "TP", "subset": "all", "value": tp},
        {"metric": "FP", "subset": "all", "value": n_r - tp},
        {"metric": "FN", "subset": "all", "value": n_h - tp_h},
        {"metric": "papers", "subset": f"annotator {args.letter}", "value": len(papers)},
        {"metric": "ResLit alleles", "subset": "all", "value": n_r},
        {"metric": "annotator genes", "subset": "all", "value": n_h},
    ])

    # Per paper (all of the annotator's papers)
    g = side.groupby("PMID")
    per_paper = pd.DataFrame({"PMID": papers}).set_index("PMID")
    per_paper["PubMed"] = "https://pubmed.ncbi.nlm.nih.gov/" + per_paper.index + "/"
    per_paper["TP"] = g.apply(lambda d: (d.status == "match").sum())
    per_paper["FP"] = g.apply(lambda d: (d.status == "only ResLit (FP)").sum())
    per_paper["FN"] = g.apply(lambda d: (d.status == "only annotator (FN)").sum())
    per_paper[["TP", "FP", "FN"]] = per_paper[["TP", "FP", "FN"]].fillna(0).astype(int)
    per_paper["differs"] = (per_paper.FP + per_paper.FN > 0).map({True: "yes", False: ""})
    per_paper["ResLit_genes"] = rl.groupby("PMID")["reslit_gene"].apply(joined)
    per_paper["ResLit_alleles"] = rl.groupby("PMID")["reslit_allele"].apply(joined)
    per_paper["annotator_genes"] = hum.groupby("PMID")["annotator_gene"].apply(joined)
    per_paper["matched"] = side[side.status == "match"].groupby("PMID")["reslit_allele"].apply(joined)
    per_paper["only_ResLit"] = side[side.status == "only ResLit (FP)"].groupby("PMID")["reslit_allele"].apply(joined)
    per_paper["only_annotator"] = side[side.status == "only annotator (FN)"].groupby("PMID")["annotator_gene"].apply(joined)
    per_paper["annotator_no_determinant_flag"] = per_paper.index.isin(no_det)
    per_paper["annotator_no_determinant_flag"] = per_paper["annotator_no_determinant_flag"].map({True: "Y", False: ""})
    per_paper = per_paper.fillna("").reset_index()

    method = pd.DataFrame({"method": [
        f"Annotator {args.letter} ({Path(args.annotator).name}) vs ResLit ({Path(args.reslit).name}), type = gene only.",
        f"Papers: the {len(papers)} papers in the annotator's file (assigned_to {args.letter} or BOTH).",
        "Annotator `gene` column compared with ResLit `allele` column, within each paper.",
        "Normalisation: trailing _<number> removed (blaTEM-1B_1 -> blaTEM-1B), lower case, all non-alphanumeric characters removed.",
        "ResLit alleles listed together ('a | b') are split; duplicates within a paper are counted once.",
        "Aliases: annotator 'X (Y)' matches on X, Y or the full name.",
        "TP = ResLit allele matched; FP = ResLit allele not matched; FN = annotator gene not matched.",
        "Precision = TP / ResLit alleles; recall = matched annotator genes / annotator genes; 95% Wilson intervals.",
    ]})

    write_workbook(out, {"summary": summary, "per_paper": per_paper, "side_by_side": side, "method": method})

    # Hand-check file: only papers with a difference. For each unmatched row,
    # the other side's names in the same paper that look alike are listed
    # (one starts with the other, or they share the first 4 characters; "bla" ignored).
    def similar(a, b):
        a, b = re.sub(r"^bla(?=.)", "", a), re.sub(r"^bla(?=.)", "", b)
        return a.startswith(b) or b.startswith(a) or (len(a) >= 4 and a[:4] == b[:4])

    diff = set(per_paper.loc[per_paper.differs == "yes", "PMID"])
    chk = side[side.PMID.isin(diff)].copy()
    cand = []
    for r in chk.itertuples():
        if r.status == "only ResLit (FP)":
            h = hum[hum.PMID == r.PMID]
            c = h.annotator_gene[h.aliases.map(lambda al: any(similar(r.normalised_name, x) for x in al))]
        elif r.status == "only annotator (FN)":
            al = aliases(r.annotator_gene)
            x = rl[rl.PMID == r.PMID]
            c = x.reslit_allele[x.norm.map(lambda n: any(similar(n, a) for a in al))]
        else:
            c = pd.Series(dtype=str)
        cand.append(" | ".join(dict.fromkeys(c)))
    chk.insert(2, "other_side_similar", cand)
    chk["manual_decision"] = ""
    chk["comment"] = ""
    write_workbook(out.with_name(f"gene_differences_{args.letter}.xlsx"),
                   {"per_PMID": per_paper[per_paper.PMID.isin(diff)], "side_by_side": chk, "method": method})

    pd.set_option("display.width", 200)
    print(summary.to_string(index=False, float_format=lambda x: f"{x:.3f}"))
    print(f"\n{(per_paper.differs == 'yes').sum()} of {len(papers)} papers differ")
    print(f"Written: {out}")


if __name__ == "__main__":
    main()

"""Field-level agreement, bootstrap CIs for weighted precision, and adjudication.

Called by all_papers_compare.py after the pooled comparison; writes
results/field_and_adjudication.xlsx and leaves the other outputs untouched.

1. Field agreement, for every matched ResLit-annotator pair (TP):
   - organism: same genus + species, case-insensitive ("E. coli" = "Escherichia
     coli"; strain, serovar and subspecies words are ignored). Pairs where
     either side has no species-level name are "not comparable" and left out
     of the rate; genus-level agreement is reported alongside.
   - antibiotics: at least one antibiotic or drug class in common. Names are
     resolved with the pipeline's own normaliser (pipeline/clean_organism_resistance.py,
     load_antibiotics / resolve_antibiotic / normalize_antibiotic, with the
     pipeline's reference lists); each resolved name also gets its drug class
     from the class sections of antibiotics_names_abreviations.txt
     (a carbapenem, penicillin, cephalosporin, monobactam or beta-lactamase
     inhibitor matches the general class "beta-lactams", but two different
     beta-lactam subclasses do not match; MLS/MLSB covers macrolides,
     lincosamides and streptogramins). Pairs where either side has no
     recognised antibiotic are left out of the rate.
2. Weighted (database-wide) precision with a bootstrap 95% CI: papers are
   resampled with replacement within their sampling stratum, 10,000
   replicates, seed 2026.
3. Adjudication: every FP and FN is listed with a stable `adjudication_id` and
   empty `category` / `comment` columns. A filled-in copy saved as
   data/adjudication.xlsx (sheet "adjudication") is read back on the next run:
   its categories are carried into the new sheet, and precision and recall
   are recomputed counting "correct", "naming/notation variant" and
   "coordinate variant" as matches.
"""

import importlib.util
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

import gene_allele_compare as G
import mutation_compare as M
from gene_allele_compare import DATA, TIERS, rate, write_workbook

REF = DATA / "reference"
PIPELINE = Path(__file__).resolve().parent / "pipeline" / "clean_organism_resistance.py"

CATEGORIES = ["correct (annotator missed it)", "naming/notation variant", "coordinate variant",
              "review- or citation-derived", "extraction error", "unresolved"]
CATEGORY_HELP = [
    "ResLit's entry is right; the annotator did not list it (FP rows).",
    "Same determinant written differently (e.g. aph(2''-li1) vs aph(2'')-Ii1, inhA T8A vs T-8A).",
    "Same mutation in another numbering system (e.g. rpoB S531L E. coli = S450L M. tuberculosis).",
    "Taken from background, cited work or a review rather than the paper's own results.",
    "ResLit extracted something wrong (wrong gene, wrong change, not a resistance determinant).",
    "Cannot be decided.",
]
MATCH_CATEGORIES = set(CATEGORIES[:3])
BOOT_REPS, BOOT_SEED = 10_000, 2026
BETA_LACTAM_SUBCLASSES = {"Carbapenems", "Penicillins", "Cephalosporins", "Monobactams", "Beta-lactamase inhibitors"}
NON_CLASS_SECTIONS = {"Other/misc antibiotics", "Combination / multi-drug entries", "TB-specific drugs"}
# MLS / MLSB names cover three classes
MULTI_CLASS = {"mls": {"Macrolides", "Lincosamides", "Streptogramins"},
               "mlsb": {"Macrolides", "Lincosamides", "Streptogramins"}}


# ── antibiotics (pipeline normaliser + drug classes) ─────────────────────────

def _load_pipeline():
    spec = importlib.util.spec_from_file_location("pipeline_clean", PIPELINE)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


class Antibiotics:
    def __init__(self):
        self.P = _load_pipeline()
        names, abbrev = REF / "antibiotics_names.txt", REF / "antibiotics_names_abreviations.txt"
        self.lookup = self.P.load_antibiotics(names, abbrev)
        self.classes = {}
        section = None
        for line in abbrev.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line.startswith("#"):
                section = line.lstrip("# ").strip()
                continue
            if not line or section is None or section in NON_CLASS_SECTIONS:
                continue
            canon = self.canonical(line)
            if canon:
                cls = {section.split(" — ")[0].split(" (")[0]}
                if section in BETA_LACTAM_SUBCLASSES:
                    cls.add("super:Beta-lactams")
                self.classes.setdefault(canon, set()).update(cls)
        for name, cls in MULTI_CLASS.items():
            self.classes.setdefault(name, set()).update(cls)

    def canonical(self, text):
        hit = self.P.resolve_antibiotic(text, self.lookup)
        return self.P.normalize_antibiotic(hit) if hit else None

    def tokens(self, text):
        """Resolved antibiotic names and their drug classes in a free-text cell."""
        names = set()
        if pd.isna(text):
            return names
        for part in re.split(r"[;|,()\[\]]|\s+and\s+|\s+or\s+|\s+[—–]\s+", str(text)):
            part = part.strip(" .:")
            if not part:
                continue
            c = self.canonical(part)
            if c:
                names.add(c)
                continue
            found = [self.canonical(p.strip()) for p in part.split("/")]
            if any(found):
                names.update(f for f in found if f)
                continue
            names.update(c for c in (self.canonical(w) for w in part.split() if len(w) > 2) if c)
        out = {"drug:" + n for n in names}
        for n in names:
            out.update(c if c.startswith("super:") else "class:" + c for c in self.classes.get(n, ()))
        return out

    @staticmethod
    def common(a, b):
        """Shared drugs / classes. A beta-lactam subclass (super:Beta-lactams) matches the
        general class Beta-lactams, but two different subclasses do not match each other."""
        c = {t for t in a & b if not t.startswith("super:")}
        for x, y in ((a, b), (b, a)):
            if "class:Beta-lactams" in x and "super:Beta-lactams" in y:
                c.add("class:Beta-lactams")
        return c


# ── organisms ────────────────────────────────────────────────────────────────

def organisms(text):
    """List of (genus, species or None) from a free-text organism cell."""
    out = []
    if pd.isna(text):
        return out
    for part in re.split(r"[;|,/]", str(text)):
        words = [w.strip("()[]").lower() for w in part.split()]
        words = [w for w in words if w]
        if not words:
            continue
        genus = words[0]
        species = words[1] if len(words) > 1 else None
        if species in (None, "sp", "sp.", "spp", "spp.") or not species.isalpha():
            species = None
        out.append((genus, species))
    return out


def _genus_eq(a, b):
    if a == b:
        return True
    for x, y in ((a, b), (b, a)):        # "e." vs "escherichia"
        if len(x) <= 2 and x.endswith(".") and y.startswith(x[0]) and len(y) > 2:
            return True
    return False


def organism_agreement(r_text, h_text):
    """(species call, genus call): True / False / None (not comparable)."""
    R, H = organisms(r_text), organisms(h_text)
    if not R or not H:
        return None, None
    genus = any(_genus_eq(a[0], b[0]) for a in R for b in H)
    Rs, Hs = [o for o in R if o[1]], [o for o in H if o[1]]
    if not Rs or not Hs:
        return None, genus
    species = any(_genus_eq(a[0], b[0]) and a[1] == b[1] for a in Rs for b in Hs)
    return species, genus


# ── helpers ──────────────────────────────────────────────────────────────────

def _similar_gene(a, b):
    # same rule as the gene hand-check files
    a, b = re.sub(r"^bla(?=.)", "", a), re.sub(r"^bla(?=.)", "", b)
    return a.startswith(b) or b.startswith(a) or (len(a) >= 4 and a[:4] == b[:4])


def _join(xs):
    return " | ".join(dict.fromkeys(x for x in xs if x))


def _annotator_rows(raw, kind, h):
    """Raw annotator rows behind one aggregated entry h."""
    d = raw[(raw.PMID == h.PMID) & raw["type"].str.strip().str.lower().eq(kind)]
    if kind == "gene":
        return d[d.gene.notna() & d.gene.map(G.annotator_norm).eq(h.norm)]
    keep = []
    for i, r in d.iterrows():
        g = r.gene if pd.notna(r.gene) else "(no gene)"
        if G.norm(g) == h.gene_norm and h.key in M.change_keys(r.nucleotide_change, r.protein_change)[0]:
            keep.append(i)
    return d.loc[keep]


def _reslit_rows(res, r):
    ids = str(r.record_id).split(",")
    return res[res.record_id.isin(ids)]


def _entry_label(kind, r):
    if kind == "gene":
        return f"{r.reslit_allele} (gene {r.reslit_gene}) [{r.reslit_tier}]"
    change = " / ".join(x for x in (r.reslit_protein, r.reslit_nucleotide) if pd.notna(x))
    return f"{r.reslit_gene} {change} [{r.reslit_tier}]"


def _rmatch(kind, r, h):
    return G.gene_match(r.norm, h) if kind == "gene" else M.mutation_match(r.gene_norm, r.key, h)


def _bootstrap_weighted(pp, rng):
    """pp: per-paper frame with stratum, weight, tp, n. Returns (point, lo, hi)."""
    point = (pp.weight * pp.tp).sum() / (pp.weight * pp.n).sum()
    num = np.zeros(BOOT_REPS)
    den = np.zeros(BOOT_REPS)
    for _, s in pp.groupby("stratum", sort=True):
        wt, wn = (s.weight * s.tp).to_numpy(), (s.weight * s.n).to_numpy()
        idx = rng.integers(0, len(s), size=(BOOT_REPS, len(s)))
        num += wt[idx].sum(axis=1)
        den += wn[idx].sum(axis=1)
    ratio = num[den > 0] / den[den > 0]
    lo, hi = np.percentile(ratio, [2.5, 97.5])
    return point, lo, hi


def _parse_category(v):
    if pd.isna(v) or not str(v).strip():
        return "", True
    s = str(v).strip().lower()
    for c in CATEGORIES:
        if s == c or s.startswith(c.split(" ")[0].split("/")[0].rstrip("-")):
            return c, True
    return str(v).strip(), False


# ── main entry ───────────────────────────────────────────────────────────────

def run(ctx, res, rawA, rawB, papA, papB, sel, out, adjudication_path):
    ab = Antibiotics()
    raws = {"A": rawA, "B": rawB}
    readers = lambda pmid: [t for t, p in (("A", papA), ("B", papB)) if pmid in p]
    strata = dict(zip(sel.PMID, sel.sampling_stratum))
    weights = dict(zip(sel.PMID, sel.weight.astype(float)))

    # previous adjudication (if any)
    prev, issues = {}, []
    if adjudication_path and Path(adjudication_path).exists():
        a = pd.read_excel(adjudication_path, sheet_name="adjudication", dtype=str)
        for r in a.itertuples():
            cat, ok = _parse_category(getattr(r, "category", None))
            prev[r.adjudication_id] = (cat, getattr(r, "comment", None) if pd.notna(getattr(r, "comment", None)) else "")
            if not ok:
                issues.append({"adjudication_id": r.adjudication_id, "issue": f"unknown category '{cat}', counted as unresolved"})

    pairs, adj, summary = [], [], []
    for kind, (rl, units, hA, hB) in ctx.items():
        hs = {"A": hA, "B": hB}

        # 1. field agreement on TP pairs
        for i, r in rl.iterrows():
            rr = _reslit_rows(res, r)
            r_org, r_res = _join(rr.organism.dropna()), _join(rr.resistance.dropna())
            for tag in "AB":
                for j, h in hs[tag][hs[tag].PMID == r.PMID].iterrows():
                    if not _rmatch(kind, r, h):
                        continue
                    hr = _annotator_rows(raws[tag], kind, h)
                    h_org, h_res = _join(hr.organism.dropna()), _join(hr.resistance.dropna())
                    sp, ge = organism_agreement(r_org, h_org)
                    tr, th = ab.tokens(r_res), ab.tokens(h_res)
                    common = ab.common(tr, th)
                    pairs.append({"type": kind, "PMID": r.PMID, "annotator": tag, "reslit_tier": r.reslit_tier,
                                  "ResLit_entry": _entry_label(kind, r),
                                  "annotator_entry": h.annotator_gene if kind == "gene" else
                                  f"{h.annotator_gene} {h.key.replace('raw:', '?')}",
                                  "ResLit_organism": r_org, "annotator_organism": h_org,
                                  "organism_species_agree": {True: "yes", False: "no", None: "not comparable"}[sp],
                                  "organism_genus_agree": {True: "yes", False: "no", None: "not comparable"}[ge],
                                  "ResLit_resistance": r_res, "annotator_resistance": h_res,
                                  "ResLit_antibiotics_resolved": ", ".join(sorted(t[5:] for t in tr if t.startswith("drug:"))),
                                  "annotator_antibiotics_resolved": ", ".join(sorted(t[5:] for t in th if t.startswith("drug:"))),
                                  "antibiotics_in_common": ", ".join(sorted(t.split(":", 1)[1] + (" (class)" if t.startswith("class:") else "") for t in common)),
                                  "antibiotic_agree": "not comparable" if not tr or not th else ("yes" if common else "no")})

        # 3. adjudication rows
        for i, r in rl[~(rl.hit_A | rl.hit_B)].iterrows():
            ident = r.norm if kind == "gene" else f"{r.gene_norm}:{r.key}"
            near = []
            for tag in readers(r.PMID):
                h = hs[tag][hs[tag].PMID == r.PMID]
                if kind == "gene":
                    near += [f"{tag}: {x.annotator_gene}" for x in h.itertuples()
                             if any(_similar_gene(r.norm, a) for a in x.aliases)]
                else:
                    near += [f"{tag}: {x.annotator_gene} {x.key.replace('raw:', '?')}" for x in h.itertuples()
                             if r.gene_norm in x.gene_aliases]
            adj.append({"adjudication_id": f"{kind}:FP:{r.PMID}:{ident}", "type": kind, "error": "FP",
                        "PMID": r.PMID, "annotators": " + ".join(readers(r.PMID)), "reslit_tier": r.reslit_tier,
                        "ResLit_entry": _entry_label(kind, r), "annotator_entry": "",
                        "near_miss_other_side": _join(near), "_row": i})
        for u in units[(units.in_ResLit == 0)].itertuples():
            ra = rl[rl.PMID == u.PMID]
            names = [n for n in (u.annotator_A, u.annotator_B) if n]
            near = []
            for r in ra.itertuples():
                if kind == "gene":
                    al = set().union(*(G.aliases(x) for n in names for x in re.split(r" \| | ; ", n)))
                    if any(_similar_gene(r.norm, a) for a in al):
                        near.append(_entry_label(kind, r))
                else:
                    genes = set().union(*(M.gene_aliases(x.rsplit(" ", 1)[0]) for n in names for x in n.split(" | ")))
                    if r.gene_norm in genes:
                        near.append(_entry_label(kind, r))
            entry = " ; ".join(f"{t}: {n}" for t, n in (("A", u.annotator_A), ("B", u.annotator_B)) if n)
            adj.append({"adjudication_id": f"{kind}:FN:{u.PMID}:{u.annotator_A}|{u.annotator_B}", "type": kind,
                        "error": "FN", "PMID": u.PMID, "annotators": " + ".join(readers(u.PMID)), "reslit_tier": "",
                        "ResLit_entry": "", "annotator_entry": entry, "near_miss_other_side": _join(near), "_row": None})

    pairs = pd.DataFrame(pairs)
    adj = pd.DataFrame(adj)
    adj["category"] = adj.adjudication_id.map(lambda k: prev.get(k, ("", ""))[0])
    adj["comment"] = adj.adjudication_id.map(lambda k: prev.get(k, ("", ""))[1])
    for k in set(prev) - set(adj.adjudication_id):
        issues.append({"adjudication_id": k, "issue": "in data/adjudication.xlsx but no longer an FP/FN; ignored"})
    n_filled = int(adj.category.ne("").sum())

    # summaries
    for kind in ctx:
        p = pairs[pairs.type == kind]
        for field, col in (("organism (genus + species)", "organism_species_agree"),
                           ("organism (genus only)", "organism_genus_agree"),
                           ("antibiotics (drug or class in common)", "antibiotic_agree")):
            comp = p[p[col] != "not comparable"]
            summary.append({"section": "1 field agreement (TP pairs)", "type": kind, **rate(field, "all pairs",
                            int((comp[col] == "yes").sum()), len(comp)),
                            "note": f"{len(p)} pairs; {len(p) - len(comp)} not comparable"})
            for tag in "AB":
                c2 = comp[comp.annotator == tag]
                summary.append({"section": "1 field agreement (TP pairs)", "type": kind,
                                **rate(field, f"annotator {tag}", int((c2[col] == "yes").sum()), len(c2))})

    rng_seed = BOOT_SEED
    for kind, (rl, units, hA, hB) in ctx.items():
        a = adj[adj.type == kind]
        is_match = a.category.isin(MATCH_CATEGORIES)
        fp_fixed = set(a[(a.error == "FP") & is_match]._row.dropna().astype(int))
        tp_raw = rl.hit_A | rl.hit_B
        tp_adj = tp_raw | rl.index.isin(fp_fixed)
        ref = units[(units.listed_by_A == 1) | (units.listed_by_B == 1)]
        found_raw = int(ref.in_ResLit.sum())
        found_adj = found_raw + int(((a.error == "FN") & is_match).sum())

        for label, tp in (("raw", tp_raw), ("adjudicated", tp_adj)):
            sec = "3 raw vs adjudicated"
            summary.append({"section": sec, "type": kind, **rate("precision", f"{label}, all", int(tp.sum()), len(rl))})
            for t in TIERS:
                m = rl.reslit_tier == t
                if m.any():
                    summary.append({"section": sec, "type": kind,
                                    **rate("precision", f"{label}, {t}", int(tp[m].sum()), int(m.sum()))})
            found = found_raw if label == "raw" else found_adj
            summary.append({"section": sec, "type": kind, **rate("recall", f"{label}, all", found, len(ref))})
            pr, rc = tp.sum() / len(rl), found / len(ref)
            summary.append({"section": sec, "type": kind, "metric": "F1", "subset": f"{label}, all",
                            "value": 2 * pr * rc / (pr + rc) if pr + rc else 0.0})

            # 2. weighted precision with bootstrap CI (same seed for every estimate)
            pp = (pd.DataFrame({"PMID": rl.PMID, "tp": tp.astype(int), "n": 1})
                    .groupby("PMID", as_index=False).sum())
            allp = pd.DataFrame({"PMID": sorted(set(papA) | set(papB))})
            pp = allp.merge(pp, how="left").fillna({"tp": 0, "n": 0})
            pp["weight"], pp["stratum"] = pp.PMID.map(weights), pp.PMID.map(strata)
            point, lo, hi = _bootstrap_weighted(pp, np.random.default_rng(rng_seed))
            summary.append({"section": "2 weighted precision (bootstrap)", "type": kind, "metric": "weighted precision",
                            "subset": label, "value": point, "ci95_low": lo, "ci95_high": hi,
                            "note": f"{BOOT_REPS} replicates, papers resampled within strata, seed {BOOT_SEED}"})
        summary.append({"section": "3 raw vs adjudicated", "type": kind, "metric": "adjudicated rows",
                        "subset": "filled in / total", "value": f"{int(a.category.ne('').sum())} / {len(a)}"})
        effective = a.category.where(a.category.isin(CATEGORIES) | a.category.eq(""), "unresolved")
        for c in CATEGORIES:
            summary.append({"section": "3 raw vs adjudicated", "type": kind, "metric": "category",
                            "subset": c, "value": int((effective == c).sum())})

    summary = pd.DataFrame(summary)
    summary = summary[["section", "type", "metric", "subset", "k", "n", "value", "ci95_low", "ci95_high", "note"]]
    summary = summary.sort_values(["section", "type"], kind="stable")
    adj = adj.drop(columns="_row").sort_values(["type", "PMID", "error", "adjudication_id"])
    cats = pd.DataFrame({"category": CATEGORIES, "meaning": CATEGORY_HELP,
                         "counted as match when recomputing": ["yes", "yes", "yes", "no", "no", "no"]})
    method = pd.DataFrame({"method": [l.strip() for l in __doc__.strip().splitlines()[2:] if l.strip()] + [
        f"Adjudication input read from: {adjudication_path if prev else 'none (no data/adjudication.xlsx yet)'}; "
        f"{n_filled} rows categorised.",
        "To adjudicate: copy this file to data/adjudication.xlsx, fill in `category` (pick from the list) and "
        "`comment` on the adjudication sheet, and re-run scripts/all_papers_compare.py. Keep `adjudication_id` unchanged.",
    ]})
    sheets = {"summary": summary, "field_pairs": pairs, "adjudication": adj, "categories": cats, "method": method}
    if issues:
        sheets["adjudication_issues"] = pd.DataFrame(issues)
    write_workbook(out, sheets, color_sheet=None)

    wb = load_workbook(out)
    for row in wb["summary"].iter_rows(min_row=2, min_col=5, max_col=9):
        for c in row:
            c.number_format = "0.000" if c.column >= 7 and isinstance(c.value, float) else "General"
    ws = wb["adjudication"]
    col = [c.value for c in ws[1]].index("category") + 1
    dv = DataValidation(type="list", formula1="=categories!$A$2:$A$7", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(1, col).column_letter}2:{ws.cell(1, col).column_letter}{ws.max_row}")
    wb.save(out)
    return summary, len(adj), n_filled, issues

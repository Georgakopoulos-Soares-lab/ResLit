#!/usr/bin/env python3
"""ResLit vs the annotators on all 50 audited papers (pooled result).

Reference list per paper: annotator A's list for A's papers, B's list for B's
papers, and the union of both lists for the 10 shared papers. Genes and
mutations are parsed and matched as in gene_allele_compare.py and
mutation_compare.py; entries are linked into units as in shared_compare.py.

Reported, separately for genes and mutations:
  precision           ResLit entries matched by the paper's annotator(s), overall and by tier
  weighted precision  the same, with each paper weighted by its sampling weight
                      (data/selected_papers.tsv), i.e. a database-wide estimate
                      that corrects for the over-sampling of Candidate papers
  recall              annotator units that ResLit also has

Usage:
    python3 scripts/all_papers_compare.py
"""

import argparse

import pandas as pd

import gene_allele_compare as G
import mutation_compare as M
from gene_allele_compare import DATA, RESULTS, TIERS, rate, read_annotator, write_workbook
from shared_compare import UNIT_FILLS, build
import field_adjudication


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--annotator-a", default=DATA / "annotator_A.xlsx")
    ap.add_argument("--annotator-b", default=DATA / "annotator_B.xlsx")
    ap.add_argument("--reslit", default=DATA / "reslit_records.tsv")
    ap.add_argument("--papers", default=DATA / "selected_papers.tsv")
    ap.add_argument("--out", default=RESULTS / "all_papers_results.xlsx")
    ap.add_argument("--field-out", default=RESULTS / "field_and_adjudication.xlsx")
    ap.add_argument("--adjudication", default=DATA / "adjudication.xlsx",
                    help="filled-in adjudication sheet (read if it exists)")
    args = ap.parse_args()

    rawA, papA, noA = read_annotator(args.annotator_a)
    rawB, papB, noB = read_annotator(args.annotator_b)
    sel = pd.read_csv(args.papers, sep="\t", dtype={"PMID": str})
    weight = dict(zip(sel.PMID, sel.weight.astype(float)))
    papers = sorted(set(papA) | set(papB))
    res = pd.read_csv(args.reslit, sep="\t", dtype=str)
    res = res[res.PMID.isin(papers)]

    summary, all_units, per_paper = [], [], []
    ctx = {}
    for kind in ("gene", "mutation"):
        if kind == "gene":
            hA, hB = G.annotator_genes(rawA), G.annotator_genes(rawB)
            rl = G.reslit_genes(res[res.type == "gene"])
        else:
            hA, hB = M.annotator_mutations(rawA), M.annotator_mutations(rawB)
            rl = M.reslit_mutations(res[res.type == "mutation"])
        hA, hB, rl = (d.reset_index(drop=True) for d in (hA, hB, rl))
        rl, units = build(kind, rl, hA, hB)
        ctx[kind] = (rl, units.copy(), hA, hB)

        tp = rl.hit_A | rl.hit_B
        ref = units[(units.listed_by_A == 1) | (units.listed_by_B == 1)]
        prec = rate("precision", "all", int(tp.sum()), len(rl))
        rec = rate("recall", "all", int(ref.in_ResLit.sum()), len(ref))
        w = rl.PMID.map(weight)
        rows = [prec,
                *[rate("precision", t, int(tp[rl.reslit_tier == t].sum()), int((rl.reslit_tier == t).sum()))
                  for t in TIERS if (rl.reslit_tier == t).any()],
                {"metric": "weighted precision", "subset": "all (sampling weights)",
                 "value": float((w * tp).sum() / w.sum())},
                rec,
                {"metric": "F1", "subset": "all",
                 "value": 2 * prec["value"] * rec["value"] / (prec["value"] + rec["value"])},
                {"metric": "TP", "subset": "ResLit entries matched", "value": int(tp.sum())},
                {"metric": "FP", "subset": "ResLit entries not matched", "value": int((~tp).sum())},
                {"metric": "FN", "subset": "annotator units ResLit missed", "value": int((ref.in_ResLit == 0).sum())},
                {"metric": "papers", "subset": "A + B (10 shared)", "value": len(papers)}]
        for r in rows:
            r["type"] = kind
        summary += rows

        units.insert(1, "type", kind)
        all_units.append(units)
        for pmid in papers:
            u = units[units.PMID == pmid]
            r_p = rl[rl.PMID == pmid]
            per_paper.append({"PMID": pmid, "type": kind,
                              "annotator": "A + B" if pmid in papA and pmid in papB else ("A" if pmid in papA else "B"),
                              "weight": weight.get(pmid),
                              "ResLit": len(r_p), "TP": int((r_p.hit_A | r_p.hit_B).sum()),
                              "FP": int((~(r_p.hit_A | r_p.hit_B)).sum()),
                              "FN": int((((u.listed_by_A == 1) | (u.listed_by_B == 1)) & (u.in_ResLit == 0)).sum()),
                              "no_determinant_flag": " ".join(t for t, s in (("A", noA), ("B", noB)) if pmid in s)})

    summary = pd.DataFrame(summary)[["type", "metric", "subset", "k", "n", "value", "ci95_low", "ci95_high"]]
    units = pd.concat(all_units, ignore_index=True)
    units = units[["PMID", "status"] + [c for c in units.columns if c not in ("PMID", "status")]]
    per_paper = pd.DataFrame(per_paper)
    per_paper.insert(1, "PubMed", "https://pubmed.ncbi.nlm.nih.gov/" + per_paper.PMID + "/")
    method = pd.DataFrame({"method": [
        f"All {len(papers)} audited papers. Reference: A's list (A's papers), B's list (B's papers), union of A and B (shared papers).",
        "Genes and mutations parsed and matched as in gene_allele_compare.py and mutation_compare.py; units as in shared_compare.py.",
        "Precision = ResLit entries matched / ResLit entries; recall = annotator units ResLit has / annotator units; 95% Wilson intervals.",
        "Weighted precision: each ResLit entry weighted by its paper's sampling weight (1 / inclusion probability, data/selected_papers.tsv).",
    ]})

    write_workbook(args.out, {"summary": summary, "units": units, "per_paper": per_paper, "method": method},
                   color_sheet="units", fills=UNIT_FILLS)
    from openpyxl import load_workbook
    wb = load_workbook(args.out)
    for row in wb["summary"].iter_rows(min_row=2, min_col=6, max_col=8):
        for c in row:
            c.number_format = "0.000"
    wb.save(args.out)

    pd.set_option("display.width", 200)
    print(summary.to_string(index=False, float_format=lambda x: f"{x:.3f}"))
    print(f"\nWritten: {args.out}")

    # field-level agreement, bootstrap CIs, adjudication (separate workbook)
    fsum, n_adj, n_filled, issues = field_adjudication.run(
        ctx, res, rawA, rawB, set(papA), set(papB), sel, args.field_out, args.adjudication)
    print()
    print(fsum.drop(columns="note").to_string(index=False, float_format=lambda x: f"{x:.3f}"))
    print(f"\nAdjudication: {n_filled} of {n_adj} FP/FN rows categorised"
          + (f"; {len(issues)} issue(s), see sheet adjudication_issues" if issues else ""))
    print(f"Written: {args.field_out}")


if __name__ == "__main__":
    main()

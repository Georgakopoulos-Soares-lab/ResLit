#!/usr/bin/env python3
"""ResLit vs both annotators on the papers they both read (shared papers).

Genes and mutations are parsed and matched exactly as in gene_allele_compare.py
and mutation_compare.py.

Units: within each paper, ResLit entries, annotator A entries and annotator B
entries are linked when they match (ResLit-A, ResLit-B, and A-B with the same
rules); each connected group is one unit (determinant). For each unit we know
whether ResLit, A and B have it.

Reported, separately for genes and mutations:
  precision  ResLit entries matched by A or B (reference = union of both lists),
             and against A alone and B alone; by tier for the union
  recall     units listed by A or B that ResLit also has; and against A / B alone
  agreement  between A and B: Cohen's kappa and percent agreement over all units
             (a unit only ResLit has counts as "neither annotator listed it")

Usage:
    python3 scripts/shared_compare.py
"""

import argparse
from pathlib import Path

import pandas as pd

import gene_allele_compare as G
import mutation_compare as M
from gene_allele_compare import DATA, RESULTS, TIERS, rate, read_annotator, write_workbook

UNIT_FILLS = {"ResLit + A + B": "C6EFCE", "ResLit + A": "E2EFDA", "ResLit + B": "E2EFDA",
              "ResLit only": "FFC7CE", "A + B (ResLit missed)": "FFEB9C",
              "A only": "FFF2CC", "B only": "FFF2CC"}


class DSU:
    def __init__(self):
        self.p = {}

    def find(self, x):
        self.p.setdefault(x, x)
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        self.p[self.find(a)] = self.find(b)


def cohen_kappa(a, b):
    n = len(a)
    if n == 0:
        return float("nan"), float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    pa, pb = sum(a) / n, sum(b) / n
    pe = pa * pb + (1 - pa) * (1 - pb)
    return ((po - pe) / (1 - pe) if pe < 1 else 1.0), po


def build(kind, rl, hA, hB):
    """Link ResLit / A / B entries into units. Returns (rl with A/B hit flags, units table)."""
    if kind == "gene":
        rmatch = lambda r, h: G.gene_match(r.norm, h)
        alike = G.genes_alike
        rlabel = lambda r: f"{r.reslit_allele} [{r.reslit_tier}]"
        hlabel = lambda h: h.annotator_gene
    else:
        rmatch = lambda r, h: M.mutation_match(r.gene_norm, r.key, h)
        alike = M.mutations_alike
        rlabel = lambda r: f"{r.reslit_gene} {r.key} [{r.reslit_tier}]"
        hlabel = lambda h: f"{h.annotator_gene} {h.key.replace('raw:', '?')}"

    dsu = DSU()
    nodes = {}
    for i, r in rl.iterrows():
        nodes[("R", i)] = r
        dsu.find(("R", i))
    for tag, h in (("A", hA), ("B", hB)):
        for i, x in h.iterrows():
            nodes[(tag, i)] = x
            dsu.find((tag, i))

    rl = rl.copy()
    rl["hit_A"], rl["hit_B"] = False, False
    for i, r in rl.iterrows():
        for tag, h in (("A", hA), ("B", hB)):
            for j, x in h[h.PMID == r.PMID].iterrows():
                if rmatch(r, x):
                    rl.at[i, f"hit_{tag}"] = True
                    dsu.union(("R", i), (tag, j))
    for i, a in hA.iterrows():
        for j, b in hB[hB.PMID == a.PMID].iterrows():
            if alike(a, b):
                dsu.union(("A", i), ("B", j))

    groups = {}
    for k in nodes:
        groups.setdefault(dsu.find(k), []).append(k)
    units = []
    for members in groups.values():
        R = [nodes[k] for k in members if k[0] == "R"]
        A = [nodes[k] for k in members if k[0] == "A"]
        B = [nodes[k] for k in members if k[0] == "B"]
        has = (bool(R), bool(A), bool(B))
        status = {(1, 1, 1): "ResLit + A + B", (1, 1, 0): "ResLit + A", (1, 0, 1): "ResLit + B",
                  (1, 0, 0): "ResLit only", (0, 1, 1): "A + B (ResLit missed)",
                  (0, 1, 0): "A only", (0, 0, 1): "B only"}[tuple(map(int, has))]
        units.append({"PMID": nodes[members[0]].PMID, "status": status,
                      "ResLit": " | ".join(rlabel(r) for r in R),
                      "annotator_A": " | ".join(hlabel(h) for h in A),
                      "annotator_B": " | ".join(hlabel(h) for h in B),
                      "in_ResLit": int(has[0]), "listed_by_A": int(has[1]), "listed_by_B": int(has[2]),
                      "best_tier": min((r.reslit_tier for r in R), key=TIERS.index) if R else "",
                      "note_A": " ; ".join(h.annotator_note for h in A if h.annotator_note),
                      "note_B": " ; ".join(h.annotator_note for h in B if h.annotator_note)})
    units = pd.DataFrame(units)
    units["_o"] = units.status.map(list(UNIT_FILLS).index)
    units = units.sort_values(["PMID", "_o", "ResLit", "annotator_A", "annotator_B"]).drop(columns="_o")
    return rl, units


def summarise(kind, rl, units):
    rows = []
    hit_any = rl.hit_A | rl.hit_B
    rows.append(rate("precision", "vs A or B (union)", int(hit_any.sum()), len(rl)))
    for t in TIERS:
        m = rl.reslit_tier == t
        if m.any():
            rows.append(rate("precision", f"vs union, {t}", int(hit_any[m].sum()), int(m.sum())))
    rows.append(rate("precision", "vs A only", int(rl.hit_A.sum()), len(rl)))
    rows.append(rate("precision", "vs B only", int(rl.hit_B.sum()), len(rl)))
    ref = units[(units.listed_by_A == 1) | (units.listed_by_B == 1)]
    rows.append(rate("recall", "vs A or B (union)", int(ref.in_ResLit.sum()), len(ref)))
    for tag in "AB":
        u = units[units[f"listed_by_{tag}"] == 1]
        rows.append(rate("recall", f"vs {tag} only", int(u.in_ResLit.sum()), len(u)))
    p, r = rows[0]["value"], rows[-3]["value"]
    rows.append({"metric": "F1", "subset": "vs union", "value": 2 * p * r / (p + r) if p + r else 0.0})
    kap, po = cohen_kappa(list(units.listed_by_A), list(units.listed_by_B))
    rows.append({"metric": "Cohen kappa (A vs B)", "subset": f"{len(units)} units", "value": kap})
    rows.append({"metric": "percent agreement (A vs B)", "subset": f"{len(units)} units", "value": po})
    # positive specific agreement: 2*both / (2*both + A only + B only); does not depend on "neither" units
    both = int(((units.listed_by_A == 1) & (units.listed_by_B == 1)).sum())
    one = int((units.listed_by_A != units.listed_by_B).sum())
    rows.append({"metric": "positive specific agreement (A vs B)", "subset": f"{both} both, {one} one only",
                 "k": 2 * both, "n": 2 * both + one, "value": 2 * both / (2 * both + one) if both + one else float("nan")})
    for st in UNIT_FILLS:
        rows.append({"metric": "units", "subset": st, "value": int((units.status == st).sum())})
    for row in rows:
        row["type"] = kind
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--annotator-a", default=DATA / "annotator_A.xlsx")
    ap.add_argument("--annotator-b", default=DATA / "annotator_B.xlsx")
    ap.add_argument("--reslit", default=DATA / "reslit_records.tsv")
    ap.add_argument("--out", default=RESULTS / "shared_papers_results.xlsx")
    args = ap.parse_args()

    rawA, papA, noA = read_annotator(args.annotator_a)
    rawB, papB, noB = read_annotator(args.annotator_b)
    res = pd.read_csv(args.reslit, sep="\t", dtype=str)
    shared = sorted(set(papA) & set(papB) & set(res.loc[res.assigned_to == "BOTH", "PMID"]))
    rawA, rawB = rawA[rawA.PMID.isin(shared)], rawB[rawB.PMID.isin(shared)]
    res = res[res.PMID.isin(shared)]

    summary, all_units, per_paper = [], [], []
    for kind in ("gene", "mutation"):
        if kind == "gene":
            hA, hB = G.annotator_genes(rawA), G.annotator_genes(rawB)
            rl = G.reslit_genes(res[res.type == "gene"])
        else:
            hA, hB = M.annotator_mutations(rawA), M.annotator_mutations(rawB)
            rl = M.reslit_mutations(res[res.type == "mutation"])
        hA, hB, rl = (d.reset_index(drop=True) for d in (hA, hB, rl))
        rl, units = build(kind, rl, hA, hB)
        summary += summarise(kind, rl, units)
        units.insert(1, "type", kind)
        all_units.append(units)
        for pmid in shared:
            u = units[units.PMID == pmid]
            k, po = cohen_kappa(list(u.listed_by_A), list(u.listed_by_B))
            per_paper.append({"PMID": pmid, "type": kind, "units": len(u),
                              "ResLit": int(u.in_ResLit.sum()), "A": int(u.listed_by_A.sum()),
                              "B": int(u.listed_by_B.sum()),
                              "A_and_B": int(((u.listed_by_A == 1) & (u.listed_by_B == 1)).sum()),
                              "A_B_percent_agreement": po,
                              **{st: int((u.status == st).sum()) for st in UNIT_FILLS},
                              "A_no_determinant": "Y" if pmid in noA else "",
                              "B_no_determinant": "Y" if pmid in noB else ""})

    summary = pd.DataFrame(summary)[["type", "metric", "subset", "k", "n", "value", "ci95_low", "ci95_high"]]
    units = pd.concat(all_units, ignore_index=True)
    units = units[["PMID", "status"] + [c for c in units.columns if c not in ("PMID", "status")]]
    per_paper = pd.DataFrame(per_paper)
    per_paper.insert(1, "PubMed", "https://pubmed.ncbi.nlm.nih.gov/" + per_paper.PMID + "/")
    method = pd.DataFrame({"method": [
        f"Shared papers: the {len(shared)} papers read by both annotators: {', '.join(shared)}.",
        "Genes and mutations parsed and matched as in gene_allele_compare.py and mutation_compare.py.",
        "Units: within a paper, ResLit, A and B entries are linked when they match (ResLit-A, ResLit-B, A-B); each linked group is one unit.",
        "Precision vs union: ResLit entries matched by A or by B / all ResLit entries. Also vs A alone and vs B alone.",
        "Recall vs union: units listed by A or B that ResLit also has / units listed by A or B. Also vs A alone and vs B alone.",
        "Cohen's kappa and percent agreement between A and B: over all units, each annotator listed it (1) or not (0); "
        "units only ResLit has count as listed by neither.",
        "Positive specific agreement = 2*both / (2*both + listed by one only); it ignores 'neither' units, which this design "
        "rarely has, so it is the more informative agreement measure here (kappa is low or negative when 'neither' is rare).",
        "95% Wilson intervals.",
    ]})

    write_workbook(args.out, {"summary": summary, "units": units, "per_paper": per_paper, "method": method},
                   color_sheet="units", fills=UNIT_FILLS)
    # summary has an extra leading column, so fix the number format of value/CI columns
    from openpyxl import load_workbook
    wb = load_workbook(args.out)
    for row in wb["summary"].iter_rows(min_row=2, min_col=6, max_col=8):
        for c in row:
            c.number_format = "0.000"
    for row in wb["per_paper"].iter_rows(min_row=2):
        for c in row:
            if per_paper.columns[c.column - 1] == "A_B_percent_agreement":
                c.number_format = "0.000"
    wb.save(args.out)

    pd.set_option("display.width", 200)
    print(f"{len(shared)} shared papers: {', '.join(shared)}\n")
    print(summary.to_string(index=False, float_format=lambda x: f"{x:.3f}"))
    print(f"\nWritten: {args.out}")


if __name__ == "__main__":
    main()

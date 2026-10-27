"""Build a 200-abstract validation sample (100 YES / 100 NO) for the Stage 3
BioMistral-7B abstract screening.

Read-only on all inputs. Writes into the directory containing this script:
  validation_200.xlsx, validation_200.csv, population_counts.txt

YES definition matches the downstream passed_pmids_*.csv files:
  is_amr == "True"  OR  prediction == "True"
(17 Vista rows have the literal string "True" in the prediction field with
is_amr False; screen_abstract_vista.sh counted them as passed.)
Everything else that has a numeric PMID is NO. Repeated header lines from
batch merging are dropped.
"""
import os
import re
import sqlite3
import xml.etree.ElementTree as ET

import numpy as np
import pandas as pd

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..")  # archive root
# papers.sqlite (PubMed abstracts + raw XML, 1.4 GB) is not in the archive; the sampled
# abstracts are in validation_200.csv
PAPERS_DB = os.path.join(ROOT, "papers.sqlite")
OUT = os.path.dirname(os.path.abspath(__file__))
SEED = 2026
N_PER_CLASS = 100
FEWSHOT_PMIDS = {"10027979", "10049269", "11557503", "1522070", "12726767"}

RESULT_FILES = [
    f"{ROOT}/pmid_lists/biomistral_scores/results_first_run.csv",
    f"{ROOT}/pmid_lists/biomistral_scores/passed_cleaned_300000.csv",
    f"{ROOT}/pmid_lists/biomistral_scores/passed_cleaned_600000.csv",
    f"{ROOT}/pmid_lists/biomistral_scores/passed_cleaned_900000.csv",
]

# 1. Combine and de-duplicate
frames = []
for f in RESULT_FILES:
    d = pd.read_csv(f, dtype=str, keep_default_na=False)
    d = d[d["pmid"].str.strip().str.fullmatch(r"\d+")]
    d["pmid"] = d["pmid"].str.strip()
    frames.append(d)
allres = pd.concat(frames, ignore_index=True)
allres["yes"] = (allres["is_amr"] == "True") | (allres["prediction"] == "True")

n_rows = len(allres)
# A PMID counts as YES if any of its rows is YES (no conflicts observed)
pop = allres.groupby("pmid", as_index=False)["yes"].max()
n_total = len(pop)
n_yes = int(pop["yes"].sum())
n_no = n_total - n_yes
print(f"rows={n_rows} unique={n_total} YES={n_yes} NO={n_no}")

# 2. Exclude few-shot examples
in_pop = sorted(FEWSHOT_PMIDS & set(pop["pmid"]))
pop = pop[~pop["pmid"].isin(FEWSHOT_PMIDS)]

# 2b. Keep only PMIDs with a non-empty abstract in papers.sqlite
con = sqlite3.connect(f"file:{PAPERS_DB}?mode=ro", uri=True)
with_abs = {r[0] for r in con.execute(
    "SELECT pmid FROM papers WHERE abstract IS NOT NULL AND TRIM(abstract) != ''")}
con.close()
n_no_abs_yes = int((~pop["pmid"].isin(with_abs) & pop["yes"]).sum())
n_no_abs_no = int((~pop["pmid"].isin(with_abs) & ~pop["yes"]).sum())
pop = pop[pop["pmid"].isin(with_abs)]

# 3. Sample with fixed seed (sorted first so the draw is order-independent)
pop = pop.sort_values("pmid", key=lambda s: s.astype(np.int64)).reset_index(drop=True)
rng = np.random.default_rng(SEED)
yes_pool = pop.loc[pop["yes"], "pmid"].to_numpy()
no_pool = pop.loc[~pop["yes"], "pmid"].to_numpy()

# Stored abstracts are often truncated at the first inline XML tag (<i>, <sub>,
# ...). A paper is kept only if its stored abstract contains every
# AbstractText section of raw_xml in full. Candidates are walked in a seeded
# random order until N_PER_CLASS pass, which is a uniform random sample of the
# full-abstract papers in each pool.
def _norm(s):
    return re.sub(r"\s+", " ", s or "").strip()


def is_full_abstract(abstract, raw_xml):
    try:
        ab = ET.fromstring(raw_xml).find(".//Abstract")
    except (ET.ParseError, TypeError):
        return False
    if ab is None:
        return False
    parts = [_norm("".join(t.itertext())) for t in ab.findall("AbstractText")]
    parts = [p for p in parts if p]
    stored = _norm(abstract)
    return bool(parts) and all(p in stored for p in parts) \
        and len(stored) >= 0.97 * sum(map(len, parts))


con = sqlite3.connect(f"file:{PAPERS_DB}?mode=ro", uri=True)
truncation_stats = {}


def draw_full(pool, label):
    picked, checked = [], 0
    for pmid in rng.permutation(pool):
        checked += 1
        abstract, raw_xml = con.execute(
            "SELECT abstract, raw_xml FROM papers WHERE pmid = ?", (pmid,)).fetchone()
        if is_full_abstract(abstract, raw_xml):
            picked.append(pmid)
            if len(picked) == N_PER_CLASS:
                break
    truncation_stats[label] = (checked, checked - len(picked))
    return np.array(picked)


yes_s = draw_full(yes_pool, "YES")
no_s = draw_full(no_pool, "NO")
con.close()
sample = pd.DataFrame({
    "pmid": np.concatenate([yes_s, no_s]),
    "biomistral_prediction": ["YES"] * N_PER_CLASS + ["NO"] * N_PER_CLASS,
})

# 4. Fetch metadata (read-only connection)
con = sqlite3.connect(f"file:{PAPERS_DB}?mode=ro", uri=True)
q = ("SELECT pmid, article_title AS title, journal, abstract, mesh_terms, keywords "
     f"FROM papers WHERE pmid IN ({','.join('?' * len(sample))})")
meta = pd.read_sql_query(q, con, params=list(sample["pmid"]))
con.close()
missing = set(sample["pmid"]) - set(meta["pmid"])
if missing:
    raise SystemExit(f"PMIDs missing from papers.sqlite: {sorted(missing)}")

# 5. Shuffle and save
out = sample.merge(meta, on="pmid", how="left")
out = out.sample(frac=1, random_state=SEED).reset_index(drop=True)
out["my_judgement"] = ""
cols = ["pmid", "title", "journal", "abstract", "mesh_terms", "keywords",
        "biomistral_prediction", "my_judgement"]
out = out[cols]
out.to_csv(f"{OUT}/validation_200.csv", index=False)

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

xlsx = f"{OUT}/validation_200.xlsx"
with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
    out.to_excel(xw, index=False, sheet_name="validation")
    ws = xw.sheets["validation"]
    widths = {"pmid": 11, "title": 40, "journal": 22, "abstract": 100,
              "mesh_terms": 35, "keywords": 25, "biomistral_prediction": 12,
              "my_judgement": 14}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[c]
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = wrap
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

# 6. Population counts
with open(f"{OUT}/population_counts.txt", "w") as fh:
    fh.write(
        "Stage 3 BioMistral-7B abstract screening - validation sample population\n"
        "Sources (de-duplicated by PMID):\n"
        + "".join(f"  {f}\n" for f in RESULT_FILES)
        + f"Total screened (unique PMIDs): {n_total}\n"
        f"Predicted YES: {n_yes}\n"
        f"Predicted NO:  {n_no}\n"
        "YES rule: is_amr == True OR prediction == 'True' (matches passed_pmids_*.csv)\n"
        f"Few-shot PMIDs excluded before sampling: {', '.join(sorted(FEWSHOT_PMIDS))}"
        f" (present in population: {', '.join(in_pop) if in_pop else 'none'})\n"
        f"Excluded for missing/empty abstract: YES={n_no_abs_yes} NO={n_no_abs_no}\n"
        f"Sampling pools after exclusions: YES={len(yes_pool)} NO={len(no_pool)}\n"
        "Full-abstract rule: stored abstract must contain every AbstractText section\n"
        "  of raw_xml (stored abstracts are often truncated at inline tags)\n"
        + "".join(f"  {k}: {c} candidates checked, {t} rejected as truncated\n"
                  for k, (c, t) in truncation_stats.items())
        + f"Sample: {N_PER_CLASS} YES + {N_PER_CLASS} NO, numpy default_rng seed {SEED}\n"
        f"Row order shuffled with pandas random_state {SEED}\n"
    )
print(open(f"{OUT}/population_counts.txt").read())
